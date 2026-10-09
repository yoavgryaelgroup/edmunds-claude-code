#!/usr/bin/env python3
"""Phase 0: load the battery dataset workbook into a PostgreSQL catalog schema.

Usage:
    pip install openpyxl "psycopg[binary]"
    DATABASE_URL=postgresql://postgres:password@localhost:5432/battery_aging \
        python load_workbook.py Battery_Datasets_with_Scraped_Data.xlsx

Reads Sheet1 (dataset list), Scraped Data and Scrape Notes, and loads them into the schema named by
DB_SCHEMA (default: catalog). Sheet1 uses merged cells to group rows: a publisher block (columns A/B)
holds dataset entries (local reference, description, chemistry, dataset link), and extra rows inside an
entry list more papers. Each entry becomes one row of catalog.dataset, and every link becomes a row of
catalog.dataset_link. Every non-empty cell is also kept as-is in catalog.sheet_cell.

Running it again replaces the catalog contents with the new workbook; catalog.workbook_import keeps
a history of what was imported and when.
"""
import hashlib
import os
import re
import sys
import urllib.parse
from pathlib import Path

import openpyxl
import psycopg
from openpyxl.utils import get_column_letter
from psycopg import sql

HERE = Path(__file__).resolve().parent
URL_RE = re.compile(r"https?://[^\s<>\"']+")
DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s<>\"']+)")
SIZE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*([TGMK])i?[Bb]\b")
SIZE_UNIT = {"T": 10**12, "G": 10**9, "M": 10**6, "K": 10**3}

# Sheet1 headers -> field names. Matched on the header text, so column order may change.
SHEET1_FIELDS = {
    "sr. no": "sr_no",
    "dataset name (publisher)": "publisher",
    "dataset no (local reference)": "local_ref",
    "basic information": "basic_information",
    "objective of study": "objective",
    "additional information": "additional_information",
    "papers affiliated": "papers",
    "relevant github link": "code",
    "dataset link": "dataset_link",
    "chemistry": "chemistry",
    "form-factor": "form_factor",
    "manufacturer": "manufacturer",
    "capacity": "capacity",
    "no. of cells tested": "cells_tested",
    "something that we have found but we don't know where this fits": "unplaced_notes",
}
ASSIGNEE_HEADERS = ["Mihir", "Irina", "Sarah", "Harshita"]
# Columns whose new value starts a new dataset entry (the rest only add papers / links).
BOUNDARY_FIELDS = ["sr_no", "local_ref", "basic_information", "objective", "additional_information",
                   "dataset_link", "chemistry", "form_factor", "manufacturer", "capacity", "cells_tested"]
LINK_KIND = {"dataset_link": "dataset", "papers": "paper", "code": "code", "unplaced_notes": "other",
             "basic_information": "reference", "objective": "reference", "additional_information": "reference",
             "chemistry": "reference", "form_factor": "reference", "manufacturer": "reference",
             "capacity": "reference", "cells_tested": "reference"}

SCRAPED_FIELDS = {
    "#": "scraped_id", "source row(s) in sheet1": "source_rows", "publisher (from sheet1)": "publisher",
    "local ref (from sheet1)": "local_ref", "dataset link": "url", "scrape status": "scrape_status",
    "scraped title": "title", "chemistry": "chemistry", "form factor": "form_factor",
    "manufacturer / model": "manufacturer_model", "nominal capacity": "nominal_capacity", "voltage": "voltage",
    "no. of cells / units": "cells_units", "test temperatures": "test_temperatures",
    "tests performed": "tests_performed", "test equipment": "test_equipment",
    "data size / format": "data_size_format", "license": "license", "published": "published",
    "summary / notes": "summary",
}

HOST_PLATFORM = [
    ("data.mendeley.com", "mendeley"), ("prod-dcd-datasets-cache-zipfiles", "mendeley"),
    ("zenodo.org", "zenodo"), ("publications.rwth-aachen.de", "rwth_publications"),
    ("figshare.com", "figshare"), ("rdr.ucl.ac.uk", "figshare"), ("irr.singaporetech.edu.sg", "figshare"),
    ("data.4tu.nl", "figshare"), ("kilthub.cmu.edu", "figshare"),
    ("deepblue.lib.umich.edu", "deepblue"), ("ora.ox.ac.uk", "ora"), ("depositonce.tu-berlin.de", "dspace"),
    ("repository.tugraz.at", "invenio_rdm"), ("data.pnnl.gov", "pnnl"), ("researchdata.edu.au", "csiro"),
    ("datadryad.org", "dryad"), ("osf.io", "osf"), ("borealisdata.ca", "dataverse"), ("nasa.gov", "nasa"),
    ("calce.umd.edu", "calce"), ("data.matr.io", "matr_io"), ("batteryarchive.org", "battery_archive"),
    ("github.com", "github"), ("ieee-dataport.org", "ieee_dataport"), ("onedrive.live.com", "onedrive"),
    ("drive.google.com", "google_drive"), ("sciencedirect.com", "sciencedirect"),
]
DOI_PLATFORM = [("10.5281/zenodo", "zenodo"), ("10.1184/", "figshare"), ("10.57760/sciencedb", "sciencedb"),
                ("10.35097/", "kit_radar"), ("10.17632/", "mendeley"), ("10.6078/", "dryad")]


# ---------------------------------------------------------------- helpers

def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def find_doi(url):
    if not url:
        return None
    m = DOI_RE.search(urllib.parse.unquote(url))
    return m.group(1).rstrip(".,;)") if m else None


def platform(url):
    if not url:
        return None
    host = urllib.parse.urlparse(url).netloc.lower()
    doi = find_doi(url) if "doi.org" in host else None
    if doi:
        for prefix, name in DOI_PLATFORM:
            if doi.lower().startswith(prefix):
                return name
        return "doi"
    for needle, name in HOST_PLATFORM:
        if needle in host:
            return name
    return host or None


def size_bytes(s):
    m = SIZE_RE.search(s or "")
    return int(float(m.group(1).replace(",", ".")) * SIZE_UNIT[m.group(2).upper()]) if m else None


def clean_url(u):
    u = u.strip().rstrip(".,;")
    while u.endswith(")") and u.count(")") > u.count("("):  # keep a ')' that closes a '(' in the URL
        u = u[:-1]
    return u


def urls_in(cell):
    """(label, url, alt_url) for each link in a cell: its hyperlink plus any URLs written in its text.

    When the cell text is just a URL and the cell also has a hyperlink, both point at the same thing
    (often a DOI shown as text, linking to the publisher's page), so they become one link with the
    shown URL kept as alt_url.
    """
    label = text(cell.value)
    target = cell.hyperlink.target.strip() if cell.hyperlink and cell.hyperlink.target else None
    shown = [clean_url(u) for u in URL_RE.findall(label or "")]
    if target and len(shown) == 1 and label == shown[0]:
        return [(label, target, shown[0] if shown[0] != target else None)]
    out, seen = [], set()
    if target:
        out.append((label, target, None))
        seen.add(target)
    for u in shown:
        if u not in seen:
            seen.add(u)
            out.append((u, u, None))
    return out


class Sheet:
    """A worksheet with merged cells resolved: get(r, c) returns the cell heading the merged block."""

    def __init__(self, ws):
        self.ws = ws
        self.anchor, self.ranges = {}, {}
        for rng in ws.merged_cells.ranges:
            self.ranges[(rng.min_row, rng.min_col)] = rng.coord
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.anchor[(r, c)] = (rng.min_row, rng.min_col)
        self.headers = {c: text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)}

    def get(self, r, c):
        return self.ws.cell(*self.anchor.get((r, c), (r, c)))

    def starts(self, r, c):
        """True when row r begins a new value in column c (heads a merged block or has its own value)."""
        a = self.anchor.get((r, c))
        return a == (r, c) if a else text(self.ws.cell(r, c).value) is not None

    def cols_by_header(self, mapping):
        found = {}
        for c, h in self.headers.items():
            if h and h.strip().lower() in mapping:
                found[mapping[h.strip().lower()]] = c
        return found

    def last_row(self):
        return max((r for r in range(1, self.ws.max_row + 1)
                    if any(text(self.ws.cell(r, c).value) for c in range(1, self.ws.max_column + 1))), default=1)


# ---------------------------------------------------------------- loaders

def load_cells(cur, name, sh):
    rows = []
    for r in range(1, sh.last_row() + 1):
        for c in range(1, sh.ws.max_column + 1):
            cell = sh.ws.cell(r, c)
            v, link = text(cell.value), (cell.hyperlink.target if cell.hyperlink else None)
            if v is None and not link:
                continue
            rows.append((name, r, c, get_column_letter(c), sh.headers.get(c), v, link, sh.ranges.get((r, c))))
    with cur.copy("COPY sheet_cell (sheet, row_num, col_num, col_letter, header, value, hyperlink, merged_range) "
                  "FROM STDIN") as cp:
        for row in rows:
            cp.write_row(row)
    return len(rows)


def load_sheet1(cur, sh):
    col = sh.cols_by_header(SHEET1_FIELDS)
    missing = [f for f in ("publisher", "local_ref", "dataset_link") if f not in col]
    if missing:
        sys.exit(f"Sheet1 is missing expected columns: {missing}")
    assignee_cols = {c: h for c, h in sh.headers.items() if h in ASSIGNEE_HEADERS}
    last = sh.last_row()

    # Group rows into entries: a new entry starts where any boundary column starts a new value.
    entries = []
    for r in range(2, last + 1):
        if not any(text(sh.ws.cell(r, c).value) or (r, c) in sh.anchor for c in range(1, sh.ws.max_column + 1)):
            continue
        if not entries or any(sh.starts(r, col[f]) for f in BOUNDARY_FIELDS if f in col):
            entries.append([r])
        else:
            entries[-1].append(r)

    def field(r, f):
        return text(sh.get(r, col[f]).value) if f in col else None

    publishers, n_links = {}, 0
    for rows in entries:
        r0, r1 = rows[0], rows[-1]
        pub_cell = sh.anchor.get((r0, col["publisher"]), (r0, col["publisher"]))
        if pub_cell not in publishers:
            pub_rows = [r for (rr, cc), a in sh.anchor.items() if a == pub_cell for r in [rr]] or [r0]
            name = field(r0, "publisher") or "(no publisher given)"
            publishers[pub_cell] = cur.execute(
                "INSERT INTO publisher (sr_no, name, first_row, last_row) VALUES (%s, %s, %s, %s) RETURNING publisher_id",
                (field(r0, "sr_no"), name, min(pub_rows), max(pub_rows))).fetchone()[0]

        local_ref, listed_size = field(r0, "local_ref"), None
        if local_ref:
            parts = re.split(r"\s*size\s*[:\-]\s*", local_ref, maxsplit=1, flags=re.I)
            if len(parts) == 2:
                local_ref, listed_size = parts[0].strip() or None, parts[1].strip()
        cells = field(r0, "cells_tested")
        link_cell = sh.get(r0, col["dataset_link"])
        link_urls = urls_in(link_cell)
        durl = link_urls[0][1] if link_urls else None
        assignees = sorted({h for c, h in assignee_cols.items() for r in rows if text(sh.ws.cell(r, c).value)})
        dataset_id = cur.execute(
            """INSERT INTO dataset (publisher_id, local_ref, listed_size, listed_size_bytes, first_row, last_row,
                   basic_information, objective, additional_information, chemistry, form_factor, manufacturer,
                   capacity, cells_tested, cells_tested_num, unplaced_notes, dataset_link_label, dataset_url,
                   platform, assignees)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
               RETURNING dataset_id""",
            (publishers[pub_cell], local_ref, listed_size, size_bytes(listed_size), r0, r1,
             field(r0, "basic_information"), field(r0, "objective"), field(r0, "additional_information"),
             field(r0, "chemistry"), field(r0, "form_factor"), field(r0, "manufacturer"), field(r0, "capacity"),
             cells, int(cells) if cells and re.fullmatch(r"\d+", cells) else None,
             field(r0, "unplaced_notes"), text(link_cell.value), durl, platform(durl), assignees or None)
        ).fetchone()[0]

        # Links: each merged block once per entry; papers row by row.
        seen = set()
        for f, kind in LINK_KIND.items():
            if f not in col:
                continue
            for r in rows:
                cell = sh.get(r, col[f])
                key = (cell.row, cell.column)
                if key in seen:
                    continue
                seen.add(key)
                for label, url, alt_url in urls_in(cell):
                    cur.execute(
                        """INSERT INTO dataset_link (dataset_id, kind, sheet_row, col_letter, label, url, alt_url,
                               platform, doi)
                           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                        (dataset_id, kind, cell.row, get_column_letter(cell.column), label, url, alt_url,
                         platform(url), find_doi(url) or find_doi(alt_url)))
                    n_links += 1
    return len(publishers), len(entries), n_links


def load_scraped(cur, sh):
    col = sh.cols_by_header(SCRAPED_FIELDS)
    if "url" not in col or "scraped_id" not in col:
        sys.exit("Scraped Data sheet is missing its '#' or 'Dataset link' column")
    n = 0
    for r in range(2, sh.last_row() + 1):
        v = {f: text(sh.get(r, c).value) for f, c in col.items()}
        if not v.get("url") and not v.get("scraped_id"):
            continue
        status = v.get("scrape_status")
        rows = [int(x) for x in re.findall(r"\d+", v.get("source_rows") or "")]
        cur.execute(
            """INSERT INTO scraped_metadata (scraped_id, source_rows, publisher, local_ref, url, platform,
                   scrape_status, status_class, title, chemistry, form_factor, manufacturer_model, nominal_capacity,
                   voltage, cells_units, test_temperatures, tests_performed, test_equipment, data_size_format,
                   size_bytes, license, published, summary)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (int(v["scraped_id"]), rows or None, v.get("publisher"), v.get("local_ref"), v.get("url"),
             platform(v.get("url")), status, status.split()[0].lower() if status else None, v.get("title"),
             v.get("chemistry"), v.get("form_factor"), v.get("manufacturer_model"), v.get("nominal_capacity"),
             v.get("voltage"), v.get("cells_units"), v.get("test_temperatures"), v.get("tests_performed"),
             v.get("test_equipment"), v.get("data_size_format"), size_bytes(v.get("data_size_format")),
             v.get("license"), v.get("published"), v.get("summary")))
        n += 1
    # Match each Sheet1 entry to its scraped row by URL, falling back to the scraped row's source rows.
    cur.execute("""UPDATE dataset d SET scraped_id = s.scraped_id FROM scraped_metadata s
                   WHERE lower(rtrim(d.dataset_url, '/')) = lower(rtrim(s.url, '/'))""")
    cur.execute("""UPDATE dataset d SET scraped_id = s.scraped_id FROM scraped_metadata s
                   WHERE d.scraped_id IS NULL AND d.first_row = ANY(s.source_rows)""")
    return n


def load_notes(cur, sh):
    section, n = None, 0
    for r in range(1, sh.last_row() + 1):
        a, b = text(sh.ws.cell(r, 1).value), text(sh.ws.cell(r, 2).value)
        if a is None:
            continue
        m = re.match(r"row (\d+)", a, re.I)
        if b is None and not m and not a.endswith("."):
            section = a  # a line with no value that isn't a sentence: section heading
            continue
        sheet_row = int(m.group(1)) if m else None
        cur.execute(
            """INSERT INTO scrape_note (section, note, value, sheet_row, dataset_id)
               VALUES (%s, %s, %s, %s,
                       (SELECT dataset_id FROM dataset WHERE %s BETWEEN first_row AND last_row LIMIT 1))""",
            (section, a, b, sheet_row, sheet_row))
        n += 1
    return n


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("set DATABASE_URL, e.g. postgresql://postgres:password@localhost:5432/battery_aging")
    path = Path(sys.argv[1])
    schema = os.environ.get("DB_SCHEMA", "catalog")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(path, data_only=True)  # data_only: formula cells give their saved values
    for needed in ("Sheet1", "Scraped Data", "Scrape Notes"):
        if needed not in wb.sheetnames:
            sys.exit(f"workbook has no sheet named {needed!r} (found {wb.sheetnames})")
    sheets = {name: Sheet(wb[name]) for name in wb.sheetnames}

    with psycopg.connect(url) as conn, conn.cursor() as cur:
        cur.execute(sql.SQL("CREATE SCHEMA IF NOT EXISTS {}").format(sql.Identifier(schema)))
        cur.execute(sql.SQL("SET search_path TO {}").format(sql.Identifier(schema)))
        # Everything but workbook_import is rebuilt from the workbook, so drop it and let schema.sql recreate
        # it: a re-run then also picks up any change to the table definitions.
        cur.execute("DROP VIEW IF EXISTS v_dataset")
        cur.execute("DROP TABLE IF EXISTS sheet_cell, scrape_note, dataset_link, dataset, publisher, "
                    "scraped_metadata CASCADE")
        cur.execute((HERE / "schema.sql").read_text(encoding="utf-8"))
        cur.execute("INSERT INTO workbook_import (file_name, sha256, sheet_names) VALUES (%s, %s, %s)",
                    (path.name, digest, wb.sheetnames))
        n_cells = sum(load_cells(cur, name, sh) for name, sh in sheets.items())
        n_pub, n_ds, n_links = load_sheet1(cur, sheets["Sheet1"])
        n_scraped = load_scraped(cur, sheets["Scraped Data"])
        n_notes = load_notes(cur, sheets["Scrape Notes"])
        unmatched = cur.execute("SELECT count(*) FROM dataset WHERE dataset_url IS NOT NULL AND scraped_id IS NULL"
                                ).fetchone()[0]
        conn.commit()

    print(f"publishers:        {n_pub}")
    print(f"dataset entries:   {n_ds}")
    print(f"links:             {n_links}")
    print(f"scraped rows:      {n_scraped}")
    print(f"scrape notes:      {n_notes}")
    print(f"cells kept as-is:  {n_cells}")
    if unmatched:
        print(f"note: {unmatched} dataset entries have a URL with no Scraped Data row")
    print(f"done: tables are in schema {schema!r}")


if __name__ == "__main__":
    main()
